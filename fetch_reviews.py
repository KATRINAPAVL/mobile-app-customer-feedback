"""Fetch Citadele Bank iOS app reviews from the LV/EE/LT App Stores and save to Excel.

Pulls customer reviews from Apple's iTunes RSS endpoint for the Citadele Bank
iOS app (ID 495139240) across the Latvian, Estonian, and Lithuanian App Stores,
deduplicates against any existing local Excel file, and writes the combined
result to ``citadele_reviews.xlsx`` in the current working directory.

Usage:
    python fetch_reviews.py
"""

from __future__ import annotations

import logging
import os
import time
from datetime import datetime
from typing import Any

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

APP_ID = "495139240"
COUNTRIES = ["lv", "ee", "lt"]
MAX_PAGES = 10
OUTPUT_FILE = "citadele_reviews.xlsx"
REQUEST_DELAY_SECONDS = 0.5
MAX_RETRIES = 3
REQUEST_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/605.1.15 (KHTML, like Gecko) Version/17.0 Safari/605.1.15"
    ),
    "Accept": "application/json",
}
COLUMNS = [
    "review_id",
    "country",
    "date",
    "rating",
    "title",
    "content",
    "author",
    "app_version",
    "fetched_at",
]

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger("citadele-reviews")


def build_url(country: str, page: int) -> str:
    return (
        f"https://itunes.apple.com/{country}/rss/customerreviews/"
        f"page={page}/id={APP_ID}/sortby=mostrecent/json"
    )


def fetch_page(country: str, page: int) -> dict[str, Any] | None:
    """Fetch one page. Returns parsed JSON, or None if the page should be skipped."""
    url = build_url(country, page)
    backoff = 1.0
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            response = requests.get(url, timeout=30, headers=REQUEST_HEADERS)
        except requests.RequestException as exc:
            log.warning("%s page %d: network error (%s), attempt %d/%d",
                        country.upper(), page, exc, attempt, MAX_RETRIES)
            if attempt == MAX_RETRIES:
                return None
            time.sleep(backoff)
            backoff *= 2
            continue

        if response.status_code == 200:
            try:
                return response.json()
            except ValueError:
                log.warning("%s page %d: invalid JSON", country.upper(), page)
                return None
        if response.status_code == 404:
            log.info("%s page %d: 404, skipping", country.upper(), page)
            return None
        if 500 <= response.status_code < 600:
            log.warning("%s page %d: %d, retry %d/%d",
                        country.upper(), page, response.status_code, attempt, MAX_RETRIES)
            if attempt == MAX_RETRIES:
                return None
            time.sleep(backoff)
            backoff *= 2
            continue
        log.warning("%s page %d: unexpected status %d, skipping",
                    country.upper(), page, response.status_code)
        return None
    return None


def parse_entries(payload: dict[str, Any], country: str, fetched_at: str) -> list[dict[str, Any]]:
    """Extract reviews from a feed payload."""
    feed = payload.get("feed", {})
    raw_entries = feed.get("entry")
    if raw_entries is None:
        return []
    if isinstance(raw_entries, dict):
        raw_entries = [raw_entries]

    reviews: list[dict[str, Any]] = []
    for entry in raw_entries:
        if "im:rating" not in entry:
            # App metadata entry that sometimes appears as the first item on page 1.
            continue
        try:
            rating = int(entry["im:rating"]["label"])
        except (KeyError, TypeError, ValueError):
            continue

        raw_date = entry.get("updated", {}).get("label", "")
        try:
            date_str = datetime.fromisoformat(raw_date.replace("Z", "+00:00")).strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        except ValueError:
            date_str = raw_date

        reviews.append({
            "review_id": entry.get("id", {}).get("label", ""),
            "country": country,
            "date": date_str,
            "rating": rating,
            "title": entry.get("title", {}).get("label", ""),
            "content": entry.get("content", {}).get("label", ""),
            "author": entry.get("author", {}).get("name", {}).get("label", ""),
            "app_version": entry.get("im:version", {}).get("label", ""),
            "fetched_at": fetched_at,
        })
    return reviews


def fetch_country(country: str, fetched_at: str) -> list[dict[str, Any]]:
    all_reviews: list[dict[str, Any]] = []
    for page in range(1, MAX_PAGES + 1):
        log.info("Fetching %s page %d/%d...", country.upper(), page, MAX_PAGES)
        payload = fetch_page(country, page)
        if payload is None:
            continue
        reviews = parse_entries(payload, country, fetched_at)
        if not reviews:
            log.info("%s page %d: no reviews returned", country.upper(), page)
        else:
            log.info("%s page %d: got %d reviews", country.upper(), page, len(reviews))
        all_reviews.extend(reviews)
        time.sleep(REQUEST_DELAY_SECONDS)
    return all_reviews


def load_existing(path: str) -> pd.DataFrame:
    if not os.path.exists(path):
        return pd.DataFrame(columns=COLUMNS)
    existing = pd.read_excel(path, dtype={"review_id": str})
    for col in COLUMNS:
        if col not in existing.columns:
            existing[col] = ""
    return existing[COLUMNS]


def write_excel(df: pd.DataFrame, path: str) -> None:
    df = df.sort_values(by="date", ascending=False, kind="stable").reset_index(drop=True)
    df.to_excel(path, index=False, engine="openpyxl")

    workbook = load_workbook(path)
    sheet = workbook.active
    bold = Font(bold=True)
    for cell in sheet[1]:
        cell.font = bold

    for col_idx, column in enumerate(df.columns, start=1):
        max_len = len(str(column))
        for value in df[column].astype(str).tolist():
            if len(value) > max_len:
                max_len = len(value)
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

    workbook.save(path)


def main() -> None:
    fetched_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    output_path = os.path.abspath(OUTPUT_FILE)

    existing = load_existing(output_path)
    existing_ids = set(existing["review_id"].astype(str).tolist())

    per_country_new: dict[str, int] = {}
    new_rows: list[dict[str, Any]] = []

    for country in COUNTRIES:
        fetched = fetch_country(country, fetched_at)
        unique_new = [r for r in fetched if r["review_id"] not in existing_ids]
        # Guard against duplicates within this run (same review on multiple pages).
        seen_in_batch: set[str] = set()
        deduped_new: list[dict[str, Any]] = []
        for review in unique_new:
            if review["review_id"] in seen_in_batch:
                continue
            seen_in_batch.add(review["review_id"])
            deduped_new.append(review)
        per_country_new[country] = len(deduped_new)
        new_rows.extend(deduped_new)
        existing_ids.update(seen_in_batch)

    if new_rows:
        combined = pd.concat([existing, pd.DataFrame(new_rows, columns=COLUMNS)], ignore_index=True)
    else:
        combined = existing.copy()

    write_excel(combined, output_path)

    per_country_total = combined.groupby("country").size().to_dict()

    print("=== Citadele Reviews Fetch Summary ===")
    for country in COUNTRIES:
        new_count = per_country_new.get(country, 0)
        total = int(per_country_total.get(country, 0))
        print(f"{country.upper()}: {new_count} new reviews (total in file: {total})")
    print("-" * 40)
    print(f"Total new reviews added: {len(new_rows)}")
    print(f"Total reviews in file: {len(combined)}")
    print(f"File saved: {output_path}")


if __name__ == "__main__":
    main()
